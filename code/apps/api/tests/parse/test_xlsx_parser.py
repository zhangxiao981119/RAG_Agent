"""xlsx_parser 测试 —— 本模块此前零覆盖（M2 起的缺口，本次随合并单元格修复补齐）。

钉住四件事：
  1. 合并单元格 fill-down：左上值 MUST 展开到区域内每一格
  2. 单元格内 \\t / \\n / \\r 消毒（TSV 结构保真）
  3. 每 sheet 一个整表 block、首行进 table_header、heading_path 为 sheet 名
  4. 空 sheet 跳过
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from app.services.parse.xlsx_parser import parse_xlsx


def _xlsx(sheets: dict[str, list[list[object]]], merges: dict[str, list[str]] | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
        for ref in (merges or {}).get(name, []):
            ws.merge_cells(ref)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_merged_cells_fill_down():
    """跨行合并的左上值 MUST 展开到区域每一格（否则延续行语义为空）。"""
    data = _xlsx(
        {"费用": [["部门", "金额"], ["研发部", 100], [None, 200]]},
        merges={"费用": ["A2:A3"]},
    )
    lines = parse_xlsx(data)[0].text.split("\n")
    assert lines[0] == "部门\t金额"
    assert lines[1] == "研发部\t100"
    assert lines[2] == "研发部\t200", "合并延续行 MUST 沿用左上值"


def test_empty_merged_range_stays_empty():
    """左上格为空的合并区没有值可填，保持空（不发明默认值）。"""
    data = _xlsx(
        {"S": [["a", "b"], [None, "x"], [None, "y"]]},
        merges={"S": ["A2:A3"]},
    )
    lines = parse_xlsx(data)[0].text.split("\n")
    assert lines[1] == "\tx"
    assert lines[2] == "\ty"


def test_cell_tab_newline_cr_sanitized():
    data = _xlsx({"S": [["名称", "值"], ["伪\t列", "多\n行\r文本"]]})
    blocks = parse_xlsx(data)
    lines = blocks[0].text.split("\n")
    assert len(lines) == 2, "换行消毒后不产生额外行"
    assert lines[1].split("\t") == ["伪 列", "多 行 文本"], "每行恰好两列"


def test_sheet_is_whole_table_block_with_header():
    data = _xlsx({"库存": [["编码", "数量"], ["M-1002", 5], ["M-1003", 8]]})
    blocks = parse_xlsx(data)
    assert len(blocks) == 1
    b = blocks[0]
    assert b.is_table is True
    assert b.table_header == ["编码", "数量"]
    assert b.heading_path == "库存"
    assert b.text == "编码\t数量\nM-1002\t5\nM-1003\t8"


def test_multi_sheet_one_block_each():
    data = _xlsx({"A": [["x", 1]], "B": [["y", 2]]})
    blocks = parse_xlsx(data)
    assert [b.heading_path for b in blocks] == ["A", "B"]


def test_empty_sheet_skipped():
    data = _xlsx({"空": [], "有": [["a", 1]]})
    blocks = parse_xlsx(data)
    assert [b.heading_path for b in blocks] == ["有"]


def test_invalid_bytes_raise_parse_error():
    import pytest

    from app.services.parse.base import ParseError

    with pytest.raises(ParseError, match="XLSX 打开失败"):
        parse_xlsx(b"not an xlsx")
