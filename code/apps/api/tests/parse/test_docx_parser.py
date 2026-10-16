"""docx_parser 测试 —— 本模块此前零覆盖。

钉住三件事（都属"看着无关紧要、改错了却会静默影响检索质量"）：

  1. **标题层级栈**：heading_path 是累积路径，遇更高级标题要正确截断
     （H1→H2→H3 后回到 H2，必须回退到 H1 之下，不能残留 H2/H3）

  2. **标题同时产出 heading_path 与独立 block** —— 这是与 markdown/pdf parser
     的**有意差异**，靠测试固定。若日后改成"标题只进 heading_path"，
     必须同步考虑：① 历史 docx 需重新索引（content 变了 → 向量不同构）
     ② chunk.to_embedding_text() 的前缀不再对 docx 冗余

  3. **表格**：整表一个 block、首行进 table_header、制表符分隔、跟随当前标题路径
"""
from __future__ import annotations

from io import BytesIO

import pytest

from app.services.parse.base import ParseError
from app.services.parse.docx_parser import parse_docx


def _docx(items, tables: int = 0, shape: tuple[int, int] = (0, 0)) -> bytes:
    """在内存里造一个 docx。

    items 元素为 (文本, 样式)：
      · 样式为 int  → 标题层级（用 add_heading，产出 "Heading N" 样式）
      · 样式为 str  → 自定义样式名
      · 样式为 None → 正文段落
    """
    from docx import Document

    doc = Document()
    for text, style in items:
        if isinstance(style, int):
            doc.add_heading(text, level=style)
        elif style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)

    rows, cols = shape
    for _ in range(tables):
        table = doc.add_table(rows=rows, cols=cols)
        for r in range(rows):
            for c in range(cols):
                table.rows[r].cells[c].text = f"r{r}c{c}"

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paths(blocks) -> dict[str, str]:
    return {b.text: b.heading_path for b in blocks}


# ── 标题层级栈 ────────────────────────────────────────────────────


def test_heading_path_is_cumulative():
    blocks = parse_docx(
        _docx(
            [
                ("制度", 1),
                ("报销", 2),
                ("一线城市 500 元", None),
                ("差旅", 3),
                ("需提前审批", None),
            ]
        )
    )

    paths = _paths(blocks)
    assert paths["一线城市 500 元"] == "制度/报销"
    assert paths["需提前审批"] == "制度/报销/差旅"


def test_heading_stack_truncates_on_higher_level():
    """从 H3 回到 H2：栈必须截断到 H1 之下，不能残留 H3。"""
    blocks = parse_docx(
        _docx(
            [
                ("A", 1),
                ("B", 2),
                ("C", 3),
                ("正文一", None),
                ("D", 2),
                ("正文二", None),
            ]
        )
    )

    paths = _paths(blocks)
    assert paths["正文一"] == "A/B/C"
    assert paths["D"] == "A/D"
    assert paths["正文二"] == "A/D"


def test_sibling_heading_does_not_accumulate():
    """连续同级标题：后者取代前者，不能拼成 A/B。"""
    blocks = parse_docx(_docx([("A", 1), ("B", 1), ("正文", None)]))

    assert _paths(blocks)["正文"] == "B"


# ── 标题作为 block（有意差异，测试固定行为） ──────────────────────


def test_heading_itself_is_also_a_block():
    """★ docx 的标题正文也进 content —— 与 markdown/pdf parser 不同。

    差异是有意的（标题词本就是典型检索词，进 content 对向量召回有利），
    代价是 chunk.to_embedding_text() 拼 [heading_path] 前缀时标题重复一次。
    """
    blocks = parse_docx(_docx([("报销标准", 1), ("一线城市 500 元", None)]))

    assert [b.text for b in blocks] == ["报销标准", "一线城市 500 元"]
    # 标题块的 heading_path 含自己（栈是先 push 再取路径）
    assert blocks[0].heading_path == "报销标准"


def test_paragraphs_are_separate_blocks():
    """每段独立成 block（不在本层合并）—— 合并交给下游 chunk._accumulate。

    「一段一 unit」把段落边界保留到 unit 层，比 markdown 那种"整块再按句子切"
    更不容易拦腰截断段落。
    """
    blocks = parse_docx(_docx([("第一段", None), ("第二段", None), ("第三段", None)]))

    assert [b.text for b in blocks] == ["第一段", "第二段", "第三段"]
    assert all(b.heading_path == "" for b in blocks)


# ── 空段落与样式兜底 ──────────────────────────────────────────────


def test_blank_paragraphs_skipped():
    blocks = parse_docx(_docx([("", None), ("   ", None), ("有内容", None), ("\t", None)]))

    assert [b.text for b in blocks] == ["有内容"]


def test_heading_style_without_number_falls_back_to_level_1():
    """样式名是 "heading" 后接非数字 → int() 失败，level 兜底为 1（不抛异常）。"""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    custom = doc.styles.add_style("HeadingCustom", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("自定义标题", style=custom)
    doc.add_paragraph("正文内容")
    buf = BytesIO()
    doc.save(buf)

    blocks = parse_docx(buf.getvalue())

    # "headingcustom".replace("heading","") = "custom" → int 失败 → level=1
    assert _paths(blocks)["自定义标题"] == "自定义标题"
    assert _paths(blocks)["正文内容"] == "自定义标题"


# ── 表格 ──────────────────────────────────────────────────────────


def test_table_is_single_block_with_header_and_heading():
    blocks = parse_docx(
        _docx([("标准", 1), ("以下是表格", None)], tables=1, shape=(3, 2))
    )

    tables = [b for b in blocks if b.is_table]
    assert len(tables) == 1

    table = tables[0]
    assert table.table_header == ["r0c0", "r0c1"]
    assert table.text.split("\n")[0] == "r0c0\tr0c1"   # 制表符分隔
    assert "r2c0\tr2c1" in table.text
    assert table.heading_path == "标准"                  # 跟随当前标题路径


def test_table_without_heading_has_empty_path():
    blocks = parse_docx(_docx([], tables=1, shape=(2, 2)))

    tables = [b for b in blocks if b.is_table]
    assert len(tables) == 1
    assert tables[0].heading_path == ""


def test_vertical_merge_fills_down():
    """纵向合并：延续格 text 为空，MUST 沿用正上方值（fill-down）。

    Word 里"部门"列跨 3 行合并 —— 不展开的话第 2/3 行整列空，
    检索期"该行属于哪个部门"语义丢失。
    """
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    merged = table.cell(0, 0).merge(table.cell(2, 0))  # 纵向合并 3 行
    merged.text = "研发部"
    table.cell(0, 1).text = "张三"
    table.cell(1, 1).text = "李四"
    table.cell(2, 1).text = "王五"
    buf = BytesIO()
    doc.save(buf)

    tables = [b for b in parse_docx(buf.getvalue()) if b.is_table]
    assert len(tables) == 1
    lines = tables[0].text.split("\n")
    assert lines[0] == "研发部\t张三"
    assert lines[1] == "研发部\t李四", "延续行 MUST 填充合并值"
    assert lines[2] == "研发部\t王五", "延续行 MUST 填充合并值"


def test_horizontal_merge_naturally_expanded():
    """横向合并：python-docx 返回同一 tc 的重复引用，text 天然展开（探针钉住行为）。"""
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "费用"
    table.cell(0, 2).text = "金额"
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"
    table.cell(1, 2).text = "c"
    buf = BytesIO()
    doc.save(buf)

    tables = [b for b in parse_docx(buf.getvalue()) if b.is_table]
    lines = tables[0].text.split("\n")
    assert lines[0] == "费用\t费用\t金额", "合并区每格展开为同值"
    assert tables[0].table_header == ["费用", "费用", "金额"]


def test_cell_tab_and_newline_sanitized():
    """单元格内的 \\t / \\n 会破坏 TSV 列结构，MUST 在序列化前消毒。"""
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "伪\t列"      # 制表符：会制造伪列
    table.cell(1, 1).text = "多\n行"      # 换行：会把一行劈成两行
    buf = BytesIO()
    doc.save(buf)

    tables = [b for b in parse_docx(buf.getvalue()) if b.is_table]
    lines = tables[0].text.split("\n")
    assert len(lines) == 2, "换行消毒后不产生额外行"
    assert lines[1].split("\t") == ["伪 列", "多 行"], "每行恰好两列"


# ── 非法输入 ──────────────────────────────────────────────────────


def test_invalid_bytes_raise_parse_error():
    """非 docx 字节流 → ParseError（而非底层异常直接冒泡）。"""
    with pytest.raises(ParseError):
        parse_docx(b"this is definitely not a docx file")
