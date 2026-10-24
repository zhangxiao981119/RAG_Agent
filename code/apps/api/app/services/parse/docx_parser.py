"""DOCX 解析 —— python-docx 按段落/标题/表格切分。

Word 文档结构：
  · Heading 1-9 → heading_path 累积（层级路径）
  · 普通段落 → 跟随当前 heading_path
  · 表格 → is_table=True，table_header=首行，chunk 整表不切
  · 段落内嵌图 → 视觉 LLM 转录为「[图片转录]」块插回文档流
  · 表格单元格内嵌图 → 转录文本写回该单元格（与「型号/名称」同一行），
    再把整表序列化成 TSV —— 图的知识必须与它所属那一行同块
    （以上两项仅 parse_docx_async + ocr_enabled=true；同步入口维持纯文本）

注意：
  · 只支持 .docx（python-docx 不支持旧格式 .doc）
  · .doc 建议先让用户转 .docx 再上传（浏览器 accept 里只放 docx）
"""
from __future__ import annotations

from io import BytesIO

from docx.oxml.ns import qn  # type: ignore[import-untyped]

from app.services.parse.base import ParsedBlock, ParseError, cell_to_text

try:
    from docx import Document  # type: ignore[import-untyped]
except ImportError:  # pragma: no cover
    raise


def _vmerge_is_continue(cell) -> bool:
    """判断单元格是否为纵向合并的延续格（w:vMerge 无 val 或 val="continue"）。

    Word 纵向合并：起始格 val="restart"（或省略 val 但有内容），延续格的
    text 恒为空 —— 直接序列化会整列丢语义（INV-M-05 之外的结构保真缺口）。
    """
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return False
    vm = tcPr.find(qn("w:vMerge"))
    if vm is None:
        return False
    return vm.get(qn("w:val")) in (None, "continue")


def _table_to_block(
    table,
    heading_path: str,
    cell_extras: list[list[str]] | None = None,
) -> ParsedBlock:
    """将 python-docx Table 转为 ParsedBlock（整表）。

    cell_extras：与 table.rows × row.cells 同形的附加文本（单元格内嵌图的
    转录结果），拼在该单元格原有文本之后 —— 图的知识必须落在它所属那一行，
    否则「AX3000」与「外观描述」分居两块，检索命中一次拿不到完整信息。

    ★ 横向合并无需处理：python-docx 的 row.cells 对 gridSpan 合并
      返回同一 tc 的重复引用，text 天然重复展开（探针实测 ['M','M']）。
      附加文本同样按 index 对齐，因此合并列会重复展开同一段转录 —— 与
      文本的既有行为一致，且 EmbeddedImageRunner 的去重会让同一张图
      只转录一次，重复的只是文本本身。
    ★ 纵向合并：延续格 text 为空，按列跟踪上一行已序列化文本做
      fill-down（与 xlsx/xls 的左上值展开同一语义）。
      ★ 附加文本 MUST 在 fill-down 之前拼接：否则"只有图、没有文字"的格子
        会被判为空格而填上正上方的值，图片信息反而被覆盖掉。
    """
    rows_data: list[list[str]] = []
    for row_idx, row in enumerate(table.rows):
        cells = [cell_to_text(c.text) for c in row.cells]
        extras = cell_extras[row_idx] if cell_extras else []
        for idx, extra in enumerate(extras):
            if idx < len(cells) and extra:
                cells[idx] = f"{cells[idx]} {extra}".strip()
        prev = rows_data[-1] if rows_data else None
        for idx, text in enumerate(cells):
            if (
                prev is not None
                and idx < len(prev)
                and text == ""
                and _vmerge_is_continue(row.cells[idx])
            ):
                cells[idx] = prev[idx]  # 纵向合并延续格：沿用正上方已序列化文本
        rows_data.append(cells)
    if not rows_data:
        return ParsedBlock(text="", heading_path=heading_path, is_table=True)

    header = rows_data[0]
    text_lines = ["\t".join(header)]
    for row in rows_data[1:]:
        text_lines.append("\t".join(row))

    return ParsedBlock(
        text="\n".join(text_lines),
        heading_path=heading_path,
        is_table=True,
        table_header=header,
    )


def parse_docx(data: bytes) -> list[ParsedBlock]:
    """解析 .docx 文件，产出段落/标题/表格块。"""
    try:
        doc = Document(BytesIO(data))
    except Exception as exc:
        raise ParseError(f"DOCX 打开失败: {exc}") from exc

    blocks: list[ParsedBlock] = []
    heading_stack: list[str] = []  # 当前活动的标题路径栈
    # ★ 与 markdown/pdf parser 的两处差异 —— **有意为之，不是遗漏**：
    #   ① 段落：每段独立成 block（那两个是 buffer 累积、遇标题才 flush）。
    #      解析层不做合并，交给下游分块服务：chunk._accumulate 会按 token 预算
    #      重组，而「一段一 unit」把段落边界保留到了 unit 层，比"整块再按句子切"
    #      更不容易拦腰截断段落。
    #   ② 标题：标题文本**既进 heading_path、也作为一个 block 进 content**。
    #      那两个 parser 的标题只进 heading_path。保留标题正文对向量召回有利
    #      （标题词本就是典型检索词）。代价见 chunk.to_embedding_text() 的说明。
    #   （原先这里有个 `prev_item` 变量，注释写「用于合并连续段落」，
    #     但只赋值、从未被读取，按 ① 的判断也不需要 —— 已删除。）

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        style_name = (para.style.name or "").lower() if para.style else ""

        if style_name.startswith("heading"):
            try:
                level = int(style_name.replace("heading", "").strip())
            except ValueError:
                level = 1
            # 更新标题栈：截断到 level-1，再 push 新标题
            heading_stack = heading_stack[: level - 1]
            heading_stack.append(text)
            # 标题本身也作为一个 block
            blocks.append(
                ParsedBlock(
                    text=text,
                    heading_path="/".join(heading_stack),
                )
            )
        else:
            heading_path = "/".join(heading_stack)
            blocks.append(
                ParsedBlock(
                    text=text,
                    heading_path=heading_path,
                )
            )

    # 处理表格
    for table in doc.tables:
        # 表格跟随当前 heading_path
        heading_path = "/".join(heading_stack)
        block = _table_to_block(table, heading_path)
        if block.text:
            blocks.append(block)

    return blocks


def _extract_paragraph_images(para) -> list[tuple[bytes, str]]:
    """取一个段落里的内嵌图（inline 图片，非外链）。

    结构路径：w:p → w:r → w:drawing → … → a:blip，@r:embed 是图片部件的
    relationship id，经 para.part.related_parts[rid] 拿到原始字节与 MIME。
    ★ 只认 r:embed（内联）：r:link 是外链引用，Word 里图不在文档包内，
      字节无从取得 —— 外链图跳过（不做网络抓取，数据不出域约束同样适用）。
    """
    out: list[tuple[bytes, str]] = []
    for blip in para._p.findall(".//" + qn("a:blip")):  # noqa: SLF001
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        try:
            part = para.part.related_parts[rid]  # noqa: SLF001
        except KeyError:
            continue
        content_type = getattr(part, "content_type", "") or ""
        if not content_type.startswith("image/"):
            continue
        out.append((part.blob, content_type))
    return out


async def parse_docx_async(data: bytes) -> list[ParsedBlock]:
    """DOCX 异步入口：正文结构与 parse_docx 一致，另把段落内嵌图转录插回文档流。

    ★ ocr_enabled=false → 等价 parse_docx（内嵌图整体跳过，只入文本，D1a 口径）；
    ★ 段落转录块插在**所属段落块之后**，heading_path 跟随当前标题栈 ——
      图的知识长在它出现的位置上，不是文档尾巴上；
    ★ 表格单元格内图的转录文本写回**该单元格**，再建表 —— 不新增块，
      整表块的行列契约（chunk 层"整表不切"）保持不变；
    ★ 过滤/去重/上限/失败语义全部收在 EmbeddedImageRunner（口径单点）。
    """
    from app.config.settings import get_settings
    from app.services.parse.embedded_images import EmbeddedImageRunner

    if not get_settings().ocr_enabled:
        return parse_docx(data)

    try:
        doc = Document(BytesIO(data))
    except Exception as exc:
        raise ParseError(f"DOCX 打开失败: {exc}") from exc

    blocks: list[ParsedBlock] = []
    heading_stack: list[str] = []
    # (插入位置=当时的 len(blocks), 该位置的图列表, 采集时的标题路径)；
    # 倒序插入避免索引漂移
    slots: list[tuple[int, list[tuple[bytes, str]], str]] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        images = _extract_paragraph_images(para)

        style_name = (para.style.name or "").lower() if para.style else ""

        if style_name.startswith("heading"):
            try:
                level = int(style_name.replace("heading", "").strip())
            except ValueError:
                level = 1
            heading_stack = heading_stack[: level - 1]
            heading_stack.append(text)
            if text:
                blocks.append(
                    ParsedBlock(
                        text=text,
                        heading_path="/".join(heading_stack),
                    )
                )
        elif text:
            blocks.append(
                ParsedBlock(
                    text=text,
                    heading_path="/".join(heading_stack),
                )
            )

        if images:
            # 纯图空段落也是合法图位：插在当前位置之后即可
            slots.append((len(blocks), images, "/".join(heading_stack)))

    runner = EmbeddedImageRunner()
    # 正序转录（文档序串行，成本护栏计数正确），倒序插入（保持记录位置有效）
    transcribed: list[tuple[int, ParsedBlock]] = []
    for idx, images, heading_path in slots:
        for blob, mime in images:
            block = await runner.process(
                blob, mime, heading_path=heading_path, page_no=None
            )
            if block is not None:
                transcribed.append((idx, block))
    # slot 的 idx 是"当时的 len(blocks)"，即紧随其后的插入位；倒序插入互不漂移
    for idx, block in reversed(transcribed):
        blocks.insert(idx, block)

    # 表格与同步入口同口径：跟随最终 heading_stack、追加在段落之后。
    # ★ 单元格图 MUST 在建表前转录完：转录文本要拼进单元格再序列化成 TSV，
    #   否则表已经切好了，图的知识只能另起一块（丢失行内绑定）。
    for table in doc.tables:
        heading_path = "/".join(heading_stack)
        cell_extras = await _transcribe_table_cells(table, heading_path, runner)
        block = _table_to_block(table, heading_path, cell_extras)
        if block.text:
            blocks.append(block)

    runner.log_summary("docx")
    return blocks


async def _transcribe_table_cells(table, heading_path: str, runner) -> list[list[str]]:
    """表格每个单元格的内嵌图 → 转录文本矩阵（与 rows × row.cells 同形）。

    ★ 只处理单元格**直接段落**里的图；单元格内再嵌套表格的图不在 v1 范围 ——
      嵌套表本身就会破坏"整表不切"的行列语义，收益远小于复杂度。
    ★ 单元格里既有的文字保持原样，转录文本以空格拼在其后（追加顺序 =
      段落顺序 = 文档顺序）；返回的每段文本已带「[图片转录]」前缀。
    ★ 预算走 "table_cell" 独立通道（见 decisions.TABLE_EMBEDDED_IMAGE_MAX_PER_DOC），
      不与段落图共用，避免段落循环先占满额度导致图片表一张都不转。
    """
    extras: list[list[str]] = []
    for row in table.rows:
        row_extras: list[str] = []
        for cell in row.cells:
            texts: list[str] = []
            for para in cell.paragraphs:
                for blob, mime in _extract_paragraph_images(para):
                    block = await runner.process(
                        blob,
                        mime,
                        heading_path=heading_path,
                        page_no=None,
                        category="table_cell",
                    )
                    if block is not None:
                        texts.append(block.text)
            row_extras.append(" ".join(texts))
        extras.append(row_extras)
    return extras
