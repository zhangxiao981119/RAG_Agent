"""PPTX 解析 —— python-pptx 按页/标题/段落/表格切分。

幻灯片结构：
  · 每页 → heading_path = "第N页"（有标题时为 "第N页 / 标题"），page_no = N
  · 标题占位符文本：进 heading_path，也作为一个 block 进正文（与 docx 同口径，
    标题词本就是典型检索词；纯标题的分节页因此不会整页消失）
  · 普通文本框段落 → 每段一个 block（与 docx 一致：段落边界保留到 unit 层）
  · 表格 → is_table=True，table_header=首行，chunk 整表不切
  · 组合（group）内的形状递归展开
  · 页码 / 日期 / 页眉页脚占位符（"1"、日期串）跳过 —— 对知识无贡献
  · 备注（notes）不抽取：备注是演讲提示，与正文高度重复
  · 图片（仅 parse_pptx_async + ocr_enabled=true）→ 视觉 LLM 转录为
    「[图片转录]」块，插在该图所在的页与位置

★ 合并单元格由 python-pptx 自己处理：横向合并后的文本落在首格（多个被合并
  格的文字以换行连接，这里由 cell_to_text 压成单行），其余被吞并的格子返回
  空串 —— 行宽仍等于网格列数，TSV 行等宽，行列不会错位（探针实测：
  合并「型号」「外观」后 row0 = ['型号\\n外观', '', '备注']）。
  本模块无需再写一遍合并填充逻辑。
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from app.services.parse.base import ParsedBlock, ParseError, cell_to_text

try:
    from pptx import Presentation  # type: ignore[import-untyped]
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER  # type: ignore[import-untyped]
except ImportError:  # pragma: no cover
    raise


# 装饰性占位符：内容对知识无贡献（页码 "1"、日期串、页脚公司名）
_DECOR_PLACEHOLDERS = frozenset(
    {
        PP_PLACEHOLDER.DATE,
        PP_PLACEHOLDER.FOOTER,
        PP_PLACEHOLDER.HEADER,
        PP_PLACEHOLDER.SLIDE_NUMBER,
    }
)

# 标题占位符：已作为标题块产出，正文遍历时必须跳过（否则标题重复成两块）。
# ★ 不能用 `shape is slide.shapes.title` 判等：python-pptx 每次访问 shapes
#   都会新建代理对象，对象身份不稳定，判等必然失败。
_TITLE_PLACEHOLDERS = frozenset(
    {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
)


@dataclass
class _Picture:
    """页内图片：异步入口按页与标题路径转录，同步入口直接忽略。"""

    blob: bytes
    mime: str
    heading_path: str
    page_no: int


def _slide_heading(slide, page_no: int) -> str:
    """页标题路径：`第N页` 或 `第N页 / 标题`（标题内换行压成空格）。"""
    title = slide.shapes.title
    text = ""
    if title is not None and title.has_text_frame:
        text = " ".join(title.text_frame.text.split())
    return f"第{page_no}页 / {text}" if text else f"第{page_no}页"


def _iter_leaf_shapes(shapes):
    """递归展开组合形状，只产出叶子（组合本身没有可提取内容）。"""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_leaf_shapes(shape.shapes)
        else:
            yield shape


def _is_decor_placeholder(shape) -> bool:
    """页码 / 日期 / 页眉页脚占位符：跳过（内容对知识无贡献）。"""
    if not shape.is_placeholder:
        return False
    return shape.placeholder_format.type in _DECOR_PLACEHOLDERS


def _is_title_placeholder(shape) -> bool:
    """标题占位符：已作为标题块产出，正文遍历时跳过。"""
    if not shape.is_placeholder:
        return False
    return shape.placeholder_format.type in _TITLE_PLACEHOLDERS


def _table_block(table, heading_path: str, page_no: int) -> ParsedBlock:
    """python-pptx Table → ParsedBlock（整表，TSV 序列化）。"""
    rows_data = [[cell_to_text(cell.text) for cell in row.cells] for row in table.rows]
    if not rows_data:
        return ParsedBlock(text="", heading_path=heading_path, page_no=page_no, is_table=True)

    header = rows_data[0]
    text_lines = ["\t".join(header)]
    for row in rows_data[1:]:
        text_lines.append("\t".join(row))
    return ParsedBlock(
        text="\n".join(text_lines),
        heading_path=heading_path,
        page_no=page_no,
        is_table=True,
        table_header=header,
    )


def _slide_items(slide, page_no: int):
    """按文档序产出该页的内容项：ParsedBlock（标题/段落/表格）或 _Picture。

    ★ 标题占位符先产出（它总在阅读顺序最前），随后按占位符类型跳过它，
      避免标题重复成两块。
    """
    heading = _slide_heading(slide, page_no)
    title = slide.shapes.title
    if title is not None and title.has_text_frame:
        text = title.text_frame.text.strip()
        if text:
            yield ParsedBlock(text=text, heading_path=heading, page_no=page_no)

    for shape in _iter_leaf_shapes(slide.shapes):
        if _is_title_placeholder(shape) or _is_decor_placeholder(shape):
            continue
        if shape.has_table:
            block = _table_block(shape.table, heading, page_no)
            if block.text:
                yield block
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = shape.image
            yield _Picture(
                blob=image.blob,
                mime=image.content_type,
                heading_path=heading,
                page_no=page_no,
            )
            continue
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            text = para.text.strip()
            if text:
                yield ParsedBlock(text=text, heading_path=heading, page_no=page_no)


def parse_pptx(data: bytes) -> list[ParsedBlock]:
    """解析 .pptx 文件，产出标题/段落/表格块（图片由异步入口处理）。"""
    try:
        prs = Presentation(BytesIO(data))
    except Exception as exc:
        raise ParseError(f"PPTX 打开失败: {exc}") from exc

    blocks: list[ParsedBlock] = []
    for page_no, slide in enumerate(prs.slides, start=1):
        for item in _slide_items(slide, page_no):
            if isinstance(item, ParsedBlock):
                blocks.append(item)
    return blocks


async def parse_pptx_async(data: bytes) -> list[ParsedBlock]:
    """PPTX 异步入口：结构与 parse_pptx 一致，另把页内图片转录成块。

    ★ ocr_enabled=false → 等价 parse_pptx（图片整体跳过，只入文本，D1a 口径）；
    ★ 转录块插在该图出现的位置上（页内形状顺序），heading_path / page_no 跟随本页；
    ★ 过滤/去重/上限/失败语义全部收在 EmbeddedImageRunner（口径单点）。
    """
    from app.config.settings import get_settings
    from app.services.parse.embedded_images import EmbeddedImageRunner

    if not get_settings().ocr_enabled:
        return parse_pptx(data)

    try:
        prs = Presentation(BytesIO(data))
    except Exception as exc:
        raise ParseError(f"PPTX 打开失败: {exc}") from exc

    runner = EmbeddedImageRunner()
    blocks: list[ParsedBlock] = []
    for page_no, slide in enumerate(prs.slides, start=1):
        for item in _slide_items(slide, page_no):
            if isinstance(item, ParsedBlock):
                blocks.append(item)
                continue
            block = await runner.process(
                item.blob,
                item.mime,
                heading_path=item.heading_path,
                page_no=item.page_no,
            )
            if block is not None:
                blocks.append(block)

    runner.log_summary("pptx")
    return blocks
