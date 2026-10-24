"""XLSX 解析 —— openpyxl 按 sheet 整表输出（手册 §3.3.1 表格整表不切）。

每个 sheet 产出 ParsedBlock：
  · is_table = True（chunk 服务整表不切，超限时按行组切并保留表头）
  · table_header = 第一行作为列头
  · text = 表格内容序列化（列头行 + 数据行，TSV 风格）
  · heading_path = sheet 名

★ 合并单元格（fill-down）：openpyxl 的 merged_cells.ranges 只能拿到
  合并信息，被覆盖格 iter_rows 出来是 None —— 直接序列化会丢语义
  （"费用总额"跨 3 行，后两行全空）。按 ranges 把左上值精确填充到
  区域内每一格（确定性展开，无猜测）。代价是放弃 read_only 流式加载
  （ReadOnlyWorksheet 无 merged_cells 属性），上传已有 max_upload_mb 上限。
★ 内嵌图（parse_xlsx_async）：按锚点定位到单元格；锚点格已有文字 → 转录
  文本拼进该格（图的知识必须与它所属那一行同块），锚点格为空 → 独立成块
  追加在 sheet 块之后。过滤/去重/预算/失败语义全部收在 EmbeddedImageRunner。
"""
from __future__ import annotations

import io
import mimetypes

from app.services.parse.base import ParsedBlock, ParseError, cell_to_text

try:
    from openpyxl import load_workbook  # type: ignore[import-untyped]
    from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]
except ImportError:  # pragma: no cover
    raise


def _build_merged_fill(ws: Worksheet, rows: list[list]) -> None:
    """将合并单元格的左上值原地填充进 rows（fill-down，行列即 rows 坐标）。

    rows 是 iter_rows(values_only=True) 的可变二维列表；
    MergedCellRange 的 min_row/min_col 为 1 基，转换为 0 基索引。
    """
    for rng in ws.merged_cells.ranges:
        top_left = rows[rng.min_row - 1][rng.min_col - 1]
        if top_left is None:
            continue  # 空合并区没有值可填
        for r in range(rng.min_row - 1, rng.max_row):
            row = rows[r]
            for c in range(rng.min_col - 1, rng.max_col):
                row[c] = top_left


def _read_rows(ws: Worksheet) -> list[list]:
    """读取整表并展开合并单元格（同步/异步两条路径共用，口径单点）。"""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if rows:
        _build_merged_fill(ws, rows)
    return rows


def _row_to_tsv(row: list[object]) -> str:
    return "\t".join(cell_to_text(v) for v in row)


def _sheet_block(
    rows: list[list],
    title: str,
    cell_extras: dict[tuple[int, int], str] | None = None,
) -> ParsedBlock:
    """一个 sheet 的二维表 → ParsedBlock（表格块）。

    cell_extras：{(0 基行, 0 基列): 转录文本}，把内嵌图的转录结果拼进锚点格。
    ★ MUST 在 _build_merged_fill 之后拼接：合并区被填满后，锚点格的既有文字
      才是最终值；提前拼会被 fill-down 覆盖掉，图的知识凭空消失。
    """
    for (r, c), extra in (cell_extras or {}).items():
        if r < len(rows) and c < len(rows[r]):
            rows[r][c] = f"{cell_to_text(rows[r][c])} {extra}".strip()

    header = [cell_to_text(c) for c in rows[0]]
    text_lines = [_row_to_tsv(rows[0])]
    for row in rows[1:]:
        text_lines.append(_row_to_tsv(row))
    return ParsedBlock(
        text="\n".join(text_lines),
        heading_path=title,
        is_table=True,
        table_header=header,
    )


def parse_xlsx(data: bytes) -> list[ParsedBlock]:
    try:
        # read_only=False：ReadOnlyWorksheet 不带 merged_cells，合并展开无从谈起
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise ParseError(f"XLSX 打开失败: {exc}") from exc

    blocks: list[ParsedBlock] = []
    try:
        for sheet in wb.worksheets:
            rows = _read_rows(sheet)
            if not rows:
                continue
            blocks.append(_sheet_block(rows, sheet.title))
    finally:
        wb.close()

    return blocks


async def parse_xlsx_async(data: bytes) -> list[ParsedBlock]:
    """XLSX 异步入口：表格结构与 parse_xlsx 一致，另把 sheet 内嵌图转录进表格。

    ★ ocr_enabled=false → 等价 parse_xlsx（内嵌图整体跳过，只入表格，D1a 口径）；
    ★ 锚点格有文字 → 转录文本拼进该单元格（category="table_cell"：压平行列
      分隔符并加「[图片转录]」前缀），整表块的行列契约保持不变；
    ★ 锚点格为空 / 锚点越界 → 独立成块，按 (row, col) 顺序追加在 sheet 块之后
      （category="inline"：表格图保留 is_table 与表头，普通图加「[图片转录]」前缀）；
    ★ openpyxl 的锚点 _from.row / _from.col 是 0 基（探针实测 "D2" → row=1,
      col=3），与 rows 的列表下标同基，无需换算。
    """
    from app.config.settings import get_settings
    from app.services.parse.embedded_images import EmbeddedImageRunner

    if not get_settings().ocr_enabled:
        return parse_xlsx(data)

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise ParseError(f"XLSX 打开失败: {exc}") from exc

    runner = EmbeddedImageRunner()
    blocks: list[ParsedBlock] = []
    try:
        for sheet in wb.worksheets:
            rows = _read_rows(sheet)
            extras, extra_blocks = await _transcribe_sheet_images(
                sheet, rows, runner
            )
            if rows:
                blocks.append(_sheet_block(rows, sheet.title, extras))
            # 纯图 sheet（没有表格行）也要把转录结果留下，否则图等于没入库
            blocks.extend(extra_blocks)
    finally:
        wb.close()

    runner.log_summary("xlsx")
    return blocks


async def _transcribe_sheet_images(
    sheet: Worksheet, rows: list[list], runner
) -> tuple[dict[tuple[int, int], str], list[ParsedBlock]]:
    """一个 sheet 的内嵌图 → (锚点格转录, 独立块列表)，按 (row, col) 文档序处理。"""
    anchored_imgs: list[tuple[int, int, object]] = []
    for img in sheet._images:
        # OneCellAnchor / TwoCellAnchor 都带 _from；拿不到锚点的图无从定位，跳过
        marker = getattr(img.anchor, "_from", None)
        if marker is None:
            continue
        anchored_imgs.append((marker.row, marker.col, img))
    anchored_imgs.sort(key=lambda item: (item[0], item[1]))

    extras: dict[tuple[int, int], str] = {}
    blocks: list[ParsedBlock] = []
    for row, col, img in anchored_imgs:
        fmt = (getattr(img, "format", None) or "").lower()
        mime = mimetypes.guess_type(f"img.{fmt}")[0] or "application/octet-stream"
        anchored = (
            row < len(rows) and col < len(rows[row]) and bool(cell_to_text(rows[row][col]))
        )
        block = await runner.process(
            img._data(),
            mime,
            heading_path=sheet.title,
            page_no=None,
            category="table_cell" if anchored else "inline",
        )
        if block is None:
            continue
        if anchored:
            extras[(row, col)] = block.text
        else:
            blocks.append(block)
    return extras, blocks
